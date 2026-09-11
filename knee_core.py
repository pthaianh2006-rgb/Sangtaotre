"""Shared serial parsing and safe report export for the desktop applications."""

import math
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

BAUD = 115200
TARGET = 90
OUT = "knee_data.xlsx"
MAX_LINE_BYTES = 4096


def preferred_serial_ports(ports):
    """List USB adapters first while preserving the driver's order within groups."""
    return sorted(ports, key=lambda port: not (
        getattr(port, "vid", None) is not None
        or str(getattr(port, "hwid", "")).upper().startswith("USB")
    ))


def serial_connection_error(port, error):
    """Preserve the original driver error and explain access conflicts conditionally."""
    detail = str(error)
    if any(marker in detail.lower() for marker in ("permissionerror", "access is denied", "access denied")):
        return (f"Windows từ chối mở {port}. Cổng có thể đang được chương trình khác sử dụng.\n"
                "Đóng Serial Monitor/Serial Plotter trong Arduino IDE hoặc ngắt kết nối "
                "của ứng dụng khác đang dùng cùng cổng, rồi thử lại.\n\n"
                f"Chi tiết gốc: {detail}")
    return f"Không mở được {port}. Kiểm tra cáp USB, chọn đúng cổng và bấm Quét.\n\nChi tiết gốc: {detail}"


def parse_knee_line(line):
    """Return flexion in degrees; retain the firmware's 180 - knee convention."""
    if not isinstance(line, str) or line.startswith("["):
        return None
    try:
        if line.startswith("Knee:"):
            knee = float(line.split(":", 1)[1].split()[0])
        elif line.count(",") == 1:
            knee = float(line.split(",", 1)[1])
        else:
            return None
    except (ValueError, IndexError):
        return None
    if not math.isfinite(knee) or not 0 <= knee <= 200:
        return None
    return round(180 - knee, 1)


class SerialLines:
    """Keep partial reads across serial timeouts and discard oversized frames."""

    def __init__(self, max_bytes=MAX_LINE_BYTES):
        self.max_bytes = max_bytes
        self.pending = bytearray()
        self.discarding = False

    def feed(self, chunk):
        lines = []
        for part in chunk.splitlines(keepends=True):
            complete = part.endswith(b"\n")
            if not self.discarding:
                self.pending.extend(part)
                if len(self.pending) > self.max_bytes:
                    self.pending.clear()
                    self.discarding = True
            if complete:
                if not self.discarding:
                    line = self.pending.decode("utf-8", errors="replace").strip()
                    if line:
                        lines.append(line)
                self.pending.clear()
                self.discarding = False
        return lines


def default_report_path():
    """Choose an unused name instead of overwriting an earlier measurement."""
    from datetime import datetime

    base = Path.cwd() / f"knee_data_{datetime.now():%Y%m%d_%H%M%S}"
    candidate = base.with_suffix(".xlsx")
    suffix = 1
    while candidate.exists():
        candidate = base.with_name(f"{base.name}_{suffix}").with_suffix(".xlsx")
        suffix += 1
    return candidate


def _validated_rows(rows):
    validated = []
    previous_time = -1.0
    for index, row in enumerate(rows, 1):
        if len(row) != 3:
            raise ValueError(f"Mẫu {index} phải có STT, thời gian và góc gập.")
        _, timestamp, flex = row
        if isinstance(timestamp, bool) or isinstance(flex, bool):
            raise ValueError("Thời gian và góc phải là số, không phải boolean.")
        timestamp, flex = float(timestamp), float(flex)
        if not math.isfinite(timestamp) or timestamp < previous_time or timestamp < 0:
            raise ValueError(f"Thời gian không hợp lệ ở mẫu {index}.")
        if not math.isfinite(flex) or not -20 <= flex <= 180:
            raise ValueError(f"Góc gập không hợp lệ ở mẫu {index}.")
        validated.append((index, timestamp, flex))
        previous_time = timestamp
    if len(validated) > 1_048_575:
        raise ValueError("Số mẫu vượt giới hạn của một trang Excel.")
    return validated


def _build_workbook(rows, target):
    wb = Workbook()
    ws = wb.active
    ws.title = "Du lieu"
    ws.append(["STT", "Thời gian (giây)", "Góc gập (độ)", "Mục tiêu (độ)"])
    for index, timestamp, flex in rows:
        ws.append([index, timestamp, flex, target])
    last = len(rows) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{last}"
    for col, width in zip("ABCD", (8, 20, 18, 18)):
        ws.column_dimensions[col].width = width
    # Styling only the numeric format avoids four border objects per sample.
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        row[0].number_format = "0.000"
        row[1].number_format = "0.0"

    chart = LineChart()
    chart.title = "GÓC GẬP ĐẦU GỐI THEO THỜI GIAN"
    chart.style = 2
    chart.y_axis.title = "Góc gập (độ)"
    chart.x_axis.title = "Thời gian (giây)"
    chart.height, chart.width = 11, 24
    chart.add_data(Reference(ws, min_col=3, max_col=4, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=2, max_row=last))
    chart.series[0].graphicalProperties.line.solidFill = "7C6AE0"
    chart.series[0].graphicalProperties.line.width = 28000
    chart.series[1].graphicalProperties.line.solidFill = "C00000"
    chart.series[1].graphicalProperties.line.prstDash = "dash"
    chart.x_axis.tickLblSkip = max(1, len(rows) // 12)
    chart.y_axis.scaling.min = min(0, min(row[2] for row in rows))
    chart.y_axis.majorUnit = 15
    ws.add_chart(chart, "F2")

    stats = wb.create_sheet("Thong ke")
    rng = f"'Du lieu'!C2:C{last}"
    trng = f"'Du lieu'!B2:B{last}"
    data = [
        ("Chỉ số", "Giá trị", "Ý nghĩa"),
        ("Số lần đo", f"=COUNT({rng})", "Tổng số mẫu ghi được"),
        ("Thời lượng đo (giây)", f"=MAX({trng})-MIN({trng})", "Khoảng thời gian từ mẫu đầu đến mẫu cuối, gồm thời gian tạm dừng"),
        ("Tốc độ lấy mẫu (Hz)", '=IF(B3>0,(B2-1)/B3,0)', "Số khoảng lấy mẫu / thời lượng; bằng 0 nếu chưa đủ thời gian"),
        ("Góc gập trung bình (độ)", f"=AVERAGE({rng})", "Góc gập trung bình cả buổi"),
        ("ROM - Góc gập lớn nhất (độ)", f"=MAX({rng})", "Góc gập lớn nhất ghi được"),
        ("Góc gập nhỏ nhất (độ)", f"=MIN({rng})", "Gần với tư thế duỗi thẳng"),
        ("Độ lệch chuẩn (độ)", f'=IF(B2>1,STDEV({rng}),0)', "Độ lệch chuẩn mẫu; bằng 0 khi chỉ có một mẫu"),
        ("Sai số chuẩn SEM (độ)", '=IF(B2>1,B8/SQRT(B2),0)', "Sai số của giá trị trung bình"),
        ("Sai số tương đối (%)", '=IF(B5<>0,B8/ABS(B5),0)', "Độ lệch chuẩn / trị tuyệt đối trung bình; bằng 0 khi trung bình bằng 0"),
        ("Mục tiêu ROM (độ)", target, "Mục tiêu cần đạt"),
        ("% đạt mục tiêu", '=IF(B11>0,B6/B11,0)', "ROM đạt được so với mục tiêu"),
    ]
    for row in data:
        stats.append(row)
    for worksheet in (ws, stats):
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
            cell.fill = PatternFill("solid", fgColor="231D36")
            cell.alignment = Alignment(horizontal="center")
    for row in range(3, 10):
        stats.cell(row, 2).number_format = "0.0"
    stats["B10"].number_format = stats["B12"].number_format = "0.0%"
    stats["A6"].fill = stats["B6"].fill = PatternFill("solid", fgColor="FFF2CC")
    for column, width in zip("ABC", (32, 16, 78)):
        stats.column_dimensions[column].width = width
    stats.freeze_panes = "A2"
    return wb


def export_excel(rows, out=OUT, target=TARGET):
    """Atomically replace a report only after the new workbook has saved fully."""
    temporary = None
    workbook = None
    try:
        rows = _validated_rows(rows)
        if not rows:
            return False, "Chưa có dữ liệu để xuất."
        target = float(target)
        if not math.isfinite(target) or not 0 < target <= 180:
            raise ValueError("Mục tiêu góc gập phải là số dương hữu hạn.")
        destination = Path(out).expanduser().absolute()
        if destination.suffix.lower() != ".xlsx":
            raise ValueError("Tệp báo cáo phải có đuôi .xlsx.")
        workbook = _build_workbook(rows, target)
        # A temporary file on the same volume makes os.replace atomic.
        with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".xlsx", prefix=".knee-report-", delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        os.replace(temporary, destination)
        return True, f"Đã xuất {destination} — {len(rows)} mẫu."
    except (OSError, ValueError, TypeError, OverflowError) as exc:
        return False, f"Không thể xuất báo cáo; dữ liệu hiện có được giữ nguyên: {exc}"
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
