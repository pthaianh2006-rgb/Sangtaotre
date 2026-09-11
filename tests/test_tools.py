"""Portable entry points, resource cleanup and data-driven reporting."""

import argparse
import contextlib
from io import StringIO
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

import web
from tien_ich import gen_report, tim_camera


class LauncherTests(unittest.TestCase):
    def test_address_and_port_validation(self):
        self.assertEqual(web.browser_url("0.0.0.0", 5050), "http://127.0.0.1:5050")
        self.assertEqual(web.browser_url("::1", 5050), "http://[::1]:5050")
        for invalid in ("0", "65536", "invalid"):
            with self.assertRaises(argparse.ArgumentTypeError):
                web.port_number(invalid)

    def test_check_does_not_run_server_or_open_browser(self):
        with patch.object(web, "select_python", return_value=Path(web.sys.executable)), \
             patch.object(web, "check_environment", return_value=0) as check, \
             patch.object(web.runpy, "run_path") as run, \
             patch.object(web.webbrowser, "open") as browser:
            self.assertEqual(web.main(["--check"]), 0)
        check.assert_called_once_with()
        run.assert_not_called()
        browser.assert_not_called()

    def test_camera_probe_releases_failed_and_successful_devices(self):
        cameras = [MagicMock(), MagicMock()]
        cameras[0].isOpened.return_value = False
        cameras[1].isOpened.return_value = True
        cameras[1].read.return_value = (True, SimpleNamespace(shape=(480, 640, 3)))
        cv2 = SimpleNamespace(CAP_DSHOW=700, CAP_ANY=0, VideoCapture=MagicMock(side_effect=cameras))
        with patch.dict("sys.modules", {"cv2": cv2}), contextlib.redirect_stdout(StringIO()):
            self.assertEqual(tim_camera.find_cameras(2), [1])
        for camera in cameras:
            camera.release.assert_called_once_with()


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def workbook(self, rows):
        path = self.folder / "samples.xlsx"
        wb = Workbook()
        wb.active.title = "Du lieu"
        wb.active.append(["STT", "Time", "Angle"])
        for index, values in enumerate(rows, 1):
            wb.active.append([index, *values])
        wb.save(path)
        wb.close()
        return path

    def test_empty_invalid_and_out_of_order_data_fail_explicitly(self):
        for rows in ([], [(0, "invalid")], [(2, 50), (1, 60)], [(0, True)], [(0, 181)]):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                gen_report.read_samples(self.workbook(rows))

    def test_single_sample_does_not_divide_by_zero(self):
        summary = gen_report.summarize([(0, 0)])
        self.assertIsNone(summary["hz"])
        self.assertIsNone(summary["sd"])
        self.assertEqual(summary["percent"], 0)

    def test_report_and_atomic_overwrite(self):
        source = self.workbook([(0, 0), (1, 60), (2, 90)])
        output = self.folder / "result.docx"
        gen_report.build_report(source, output)
        from docx import Document
        self.assertIn("3", [row.cells[1].text for row in Document(output).tables[0].rows])
        previous = output.read_bytes()
        with self.assertRaises(FileExistsError):
            gen_report.build_report(source, output)
        with patch("docx.document.Document.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                gen_report.build_report(source, output, overwrite=True)
        self.assertEqual(output.read_bytes(), previous)
        self.assertEqual(list(self.folder.glob(".report-*")), [])


if __name__ == "__main__":
    unittest.main()
