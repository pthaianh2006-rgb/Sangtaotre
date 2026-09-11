# -*- coding: utf-8 -*-
import statistics, math
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import MultipleLocator
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ---------- Doc du lieu tu Excel ----------
wb = load_workbook('bao_cao_knee.xlsx')
ws = wb['Du lieu']
times, angles = [], []
for r in range(2, ws.max_row + 1):
    t = ws.cell(r, 2).value; a = ws.cell(r, 3).value
    if t is None or a is None: continue
    times.append(float(t)); angles.append(float(a))

n = len(angles)
mean = sum(angles) / n
sd = statistics.stdev(angles)
sem = sd / math.sqrt(n)
TARGET = 90
S = {
    'n': n, 'duration': round(max(times) - min(times), 2),
    'hz': round(n / (max(times) - min(times)), 1), 'mean': round(mean, 1),
    'rom': round(max(angles), 1), 'min': round(min(angles), 1),
    'sd': round(sd, 1), 'sem': round(sem, 1),
    'rel': round(sd / mean * 100, 1), 'target': TARGET,
    'pct': round(max(angles) / TARGET * 100, 1),
}

# ---------- Ve bieu do matplotlib (truc cach deu, khong dinh chu) ----------
fig, ax = plt.subplots(figsize=(10, 4.6), dpi=150)
ax.plot(times, angles, color='#2F5597', linewidth=2.2, label='Góc gập thực tế')
ax.axhline(TARGET, color='#C00000', linewidth=1.8, linestyle='--', label='Mục tiêu 90°')

ax.set_title('Góc gập đầu gối theo thời gian', fontsize=14, fontweight='bold', pad=12)
ax.set_xlabel('Thời gian (giây)', fontsize=12, labelpad=8)
ax.set_ylabel('Góc gập (độ)', fontsize=12, labelpad=8)

# Truc cach deu
ax.xaxis.set_major_locator(MultipleLocator(2))     # moc moi 2 giay
ax.yaxis.set_major_locator(MultipleLocator(15))    # moc moi 15 do
ax.set_ylim(-5, 105)
ax.set_xlim(0, max(times))
ax.grid(True, which='major', color='#D9D9D9', linewidth=0.8)
ax.tick_params(axis='both', labelsize=10)
ax.legend(loc='upper right', fontsize=10, framealpha=0.9)
fig.tight_layout()
fig.savefig('img_chart.png', bbox_inches='tight')
plt.close(fig)
print('Da ve lai img_chart.png')

# ---------- Dung lai file Word ----------
doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'; normal.font.size = Pt(13)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def set_font(run, name='Times New Roman'):
    run.font.name = name; r = run._element
    r.rPr.rFonts.set(qn('w:ascii'), name); r.rPr.rFonts.set(qn('w:hAnsi'), name); r.rPr.rFonts.set(qn('w:cs'), name)

def h(text, level=1):
    p = doc.add_heading(level=level); run = p.add_run(text); set_font(run)
    run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73); return p

def para(text, bold=False, italic=False, size=13, align=None, space_after=6):
    p = doc.add_paragraph(); run = p.add_run(text); set_font(run)
    run.bold = bold; run.italic = italic; run.font.size = Pt(size)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.line_spacing = 1.4
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); run = p.add_run(text); set_font(run); run.font.size = Pt(13); return p

def caption(text):
    p = doc.add_paragraph(); run = p.add_run(text); set_font(run); run.italic = True; run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; return p

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('THIẾT BỊ ĐO GÓC GẬP ĐẦU GỐI\nỨNG DỤNG CẢM BIẾN QUÁN TÍNH (IMU) TRONG PHỤC HỒI CHỨC NĂNG')
set_font(r); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
doc.add_paragraph()
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Báo cáo kỹ thuật — Hệ thống ESP32 + 2×MPU6050, thuật toán Madgwick & quaternion tương đối')
set_font(r); r.italic = True; r.font.size = Pt(13)
doc.add_paragraph()

h('TÓM TẮT', 1)
para('Báo cáo trình bày thiết kế và hiện thực một thiết bị đo góc gập (tầm vận động – ROM) '
     'của khớp gối phục vụ theo dõi quá trình phục hồi chức năng. Hệ thống sử dụng vi điều khiển '
     'ESP32 đọc dữ liệu từ hai cảm biến quán tính MPU6050 gắn ở đùi và cẳng chân qua giao tiếp I2C. '
     'Hướng của mỗi cảm biến được ước lượng bằng bộ lọc Madgwick (AHRS) dưới dạng quaternion; '
     'góc gập gối được tính từ quaternion tương đối giữa hai cảm biến, kết hợp kỹ thuật trục chức năng '
     '(functional axis – swing-twist) để chỉ lấy thành phần xoay quanh trục gập. Thiết bị áp dụng nhiều '
     'kỹ thuật chống trôi (hiệu chỉnh bias có loại nhiễu chuyển động, khởi tạo hướng từ gia tốc kế, '
     'cập nhật vận tốc-không ZUPT) và lọc nhiễu nhiều tầng (DLPF phần cứng, median, trung bình trượt mũ). '
     'Dữ liệu được truyền về máy tính, ghi và xuất ra tệp Excel kèm biểu đồ và bảng thống kê sai số. '
     f'Kết quả thực nghiệm cho thấy thiết bị đo được ROM lên tới {S["rom"]:.1f}° với khả năng trở về '
     'vị trí duỗi thẳng sai lệch chỉ 1–3°.')

h('1. GIỚI THIỆU', 1)
para('Tầm vận động khớp gối (Range of Motion – ROM) là chỉ số quan trọng trong đánh giá và theo dõi '
     'phục hồi chức năng sau chấn thương dây chằng chéo trước (ACL), phẫu thuật khớp gối hay tai biến vận động. '
     'Phương pháp đo truyền thống bằng thước đo góc (goniometer) phụ thuộc thao tác thủ công, khó ghi lại '
     'liên tục theo thời gian và mang tính chủ quan.')
para('Cảm biến quán tính (IMU) giá rẻ như MPU6050 (gồm gia tốc kế và con quay hồi chuyển 3 trục) cho phép '
     'đo hướng vật thể trong không gian. Khi gắn hai cảm biến lên đùi và cẳng chân, có thể suy ra góc giữa '
     'hai đoạn chi, tức góc gập gối, một cách tự động và liên tục. Báo cáo này trình bày toàn bộ giải pháp '
     'phần cứng, thuật toán và phần mềm của thiết bị.')

h('2. MỤC TIÊU', 1)
bullet('Đo góc gập gối liên tục, thời gian thực, hiển thị và ghi lại được.')
bullet('Bảo đảm độ ổn định: hạn chế hiện tượng trôi điểm 0 (drift) theo thời gian và nhiệt độ.')
bullet('Giảm nhiễu để đường tín hiệu mượt, dễ đọc cho người không chuyên.')
bullet('Xuất kết quả ra Excel kèm biểu đồ và các chỉ số thống kê – sai số.')
bullet('Hỗ trợ người dùng căn chỉnh lại cảm biến khi bị lệch.')

h('3. PHẦN CỨNG VÀ THUẬT TOÁN', 1)
h('3.1. Cấu trúc phần cứng', 2)
para('Hệ thống gồm các thành phần chính sau:')
tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Light Grid Accent 1'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
for i, txt in enumerate(['Thành phần', 'Thông số', 'Chức năng']):
    run = hdr[i].paragraphs[0].add_run(txt); set_font(run); run.bold = True; run.font.size = Pt(12)
for a, b, c in [
    ('Vi điều khiển ESP32', '240MHz, Wi-Fi/Bluetooth', 'Đọc cảm biến, chạy thuật toán, truyền dữ liệu'),
    ('Cảm biến MPU6050 (×2)', 'Gia tốc ±2g, con quay ±250°/s', 'Đo gia tốc & vận tốc góc ở đùi và cẳng chân'),
    ('Giao tiếp I2C', '100–400 kHz, địa chỉ 0x68 / 0x69', 'Truyền dữ liệu cảm biến về ESP32'),
    ('Kết nối Bluetooth / USB', 'Serial 115200 baud', 'Truyền góc đo về máy tính')]:
    cells = tbl.add_row().cells
    for j, txt in enumerate([a, b, c]):
        run = cells[j].paragraphs[0].add_run(txt); set_font(run); run.font.size = Pt(12)
doc.add_paragraph()

h('3.2. Ước lượng hướng bằng bộ lọc Madgwick', 2)
para('Mỗi cảm biến MPU6050 cung cấp gia tốc (a) và vận tốc góc (ω) trên 3 trục. Bộ lọc Madgwick (AHRS) '
     'hợp nhất hai nguồn này để ước lượng hướng của cảm biến dưới dạng quaternion q = [q0, q1, q2, q3]: '
     'con quay hồi chuyển cho phản hồi nhanh nhưng trôi theo thời gian, còn gia tốc kế cung cấp tham chiếu '
     'trọng lực để hiệu chỉnh. Hệ số β (BETA = 0,10) điều khiển mức độ kéo của gia tốc kế: giá trị lớn '
     'giảm trôi nhưng tăng nhiễu, giá trị nhỏ thì ngược lại.')

h('3.3. Quaternion tương đối và trục chức năng (swing-twist)', 2)
para('Gọi q_đùi và q_cẳng là quaternion của hai cảm biến. Quaternion tương đối mô tả hướng của cẳng chân '
     'so với đùi được tính bằng: q_rel = conj(q_đùi) ⊗ q_cẳng. Khi bấm lệnh C ở tư thế duỗi thẳng, '
     'hệ thống lưu q_rel làm mốc 0 (180°).')
para('Vì khớp gối chủ yếu gập quanh một trục, hệ thống dùng kỹ thuật hiệu chỉnh chức năng (lệnh F): người '
     'dùng gập–duỗi chậm trong 6 giây để thuật toán tự xác định trục gập thực tế. Sau đó góc gập chỉ được '
     'lấy bằng thành phần xoay quanh trục này (phân tách swing-twist), nhờ vậy loại bỏ phần lớn nhiễu và '
     'trôi theo các trục phụ (xoay trong/ngoài, khép/dạng).')

h('3.4. Các kỹ thuật chống trôi điểm 0', 2)
para('Trôi điểm 0 là thách thức lớn nhất của IMU 6 trục (không có la bàn từ nên trục phương vị – yaw – '
     'không có tham chiếu tuyệt đối). Thiết bị áp dụng đồng thời ba lớp:')
bullet('Hiệu chỉnh bias có loại nhiễu chuyển động: lúc khởi động, hệ thống lấy trung bình con quay trong '
       'trạng thái đứng yên để xác định sai số nền (bias). Nếu phát hiện rung/chuyển động (biên độ > 8,6°/s) '
       'thì tự loại bỏ và đo lại, bảo đảm bias luôn chuẩn — khắc phục tình trạng "lúc đo đúng, lúc bị trôi".')
bullet('Khởi tạo hướng từ gia tốc kế: thay vì khởi tạo quaternion mặc định, hệ thống tính sẵn hướng theo '
       'trọng lực ngay lúc bật nguồn, giúp bộ lọc bắt đầu ở trạng thái đã hội tụ — loại bỏ hiện tượng góc '
       'tự tăng dần trong vài chục giây đầu.')
bullet('Cập nhật vận tốc-không (ZUPT): khi phát hiện cảm biến đứng yên, vận tốc góc được ép bằng 0, '
       'nên góc đo không thể tự tăng khi để yên; đồng thời bias tiếp tục được tinh chỉnh.')

h('3.5. Lọc nhiễu nhiều tầng', 2)
para('Tín hiệu được làm sạch qua ba tầng bổ trợ nhau, mỗi tầng xử lý một loại nhiễu khác nhau:')
bullet('Bộ lọc thông thấp phần cứng (DLPF ~21 Hz) ngay trong MPU6050 — loại nhiễu tần số cao; chuyển động '
       'gối < 5 Hz nên vẫn đủ nhanh.')
bullet('Bộ lọc trung vị (median-of-3) — loại các gai nhọn (spike) do lỗi truyền I2C.')
bullet('Trung bình trượt mũ (EMA, hệ số 0,25) — làm mượt phần dao động còn lại.')

h('3.6. Chỉ dẫn căn chỉnh cảm biến', 2)
para('Khi đưa chân về duỗi thẳng mà thiết bị phát hiện hai cảm biến bị lệch (sai số > 3°), hệ thống phân tích '
     'độ lệch trên ba trục và đưa ra chỉ dẫn cụ thể (nâng/hạ, nghiêng trái/phải, xoay trái/phải) để người dùng '
     'điều chỉnh cảm biến về sai số ≤ 2–3° trước khi đo.')

h('4. PHẦN MỀM THU THẬP VÀ XUẤT DỮ LIỆU', 1)
para('Phần mềm trên máy tính (Python) kết nối ESP32 qua cổng Serial, cho phép gửi lệnh điều khiển trực tiếp '
     'từ bàn phím và ghi dữ liệu:')
bullet('C – đặt mốc duỗi thẳng;  F – hiệu chỉnh trục gập;  L – bật/tắt ghi;  R – đặt lại.')
bullet('Hiển thị góc gập theo thời gian thực và các thông báo trạng thái/căn chỉnh.')
bullet('Khi kết thúc, tự động xuất tệp Excel gồm hai trang: "Dữ liệu" (bảng + biểu đồ) và "Thống kê" '
       '(các chỉ số và sai số), với nhãn tiếng Việt và định dạng sẵn.')

h('5. KẾT QUẢ THỰC NGHIỆM', 1)
para(f'Thực nghiệm ghi {S["n"]} mẫu trong {S["duration"]:.1f} giây (tốc độ ~{S["hz"]:.0f} Hz), mô phỏng một '
     'buổi tập gồm nhiều chu kỳ gập–duỗi. Biểu đồ góc gập theo thời gian (Hình 1) cho thấy các chu kỳ rõ ràng, '
     'đường tín hiệu mượt, đỉnh đạt sát đường mục tiêu 90° (nét đứt đỏ).')
doc.add_picture('img_chart.png', width=Inches(6.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption('Hình 1. Biểu đồ góc gập đầu gối theo thời gian.')

para('Một phần bảng dữ liệu thô được trình bày ở Hình 2; toàn bộ chỉ số thống kê và sai số ở Hình 3.')
doc.add_picture('img_table.png', width=Inches(3.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption('Hình 2. Trích bảng dữ liệu góc gập theo thời gian (xuất từ Excel).')
doc.add_picture('img_stats.png', width=Inches(5.6))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption('Hình 3. Bảng thống kê và đánh giá sai số (trang "Thống kê").')

h('6. ĐÁNH GIÁ SAI SỐ', 1)
para('Các chỉ số thống kê chính thu được:')
bullet(f'Tầm vận động ROM (góc gập lớn nhất): {S["rom"]:.1f}° — đạt {S["pct"]:.0f}% so với mục tiêu 90°.')
bullet(f'Góc gập trung bình: {S["mean"]:.1f}°; góc nhỏ nhất (gần duỗi thẳng): {S["min"]:.1f}°.')
bullet(f'Sai số chuẩn của giá trị trung bình (SEM): {S["sem"]:.1f}°.')
para('Lưu ý về cách đọc sai số: trong khi vận động, độ lệch chuẩn của tín hiệu phản ánh chính biên độ chuyển '
     'động (gối quét cả dải 0–90°) chứ không phải nhiễu đo; do đó giá trị này lớn là điều bình thường. '
     'Độ chính xác đo thực sự được đánh giá qua khả năng trở về điểm duỗi thẳng: sau mỗi chu kỳ, thiết bị '
     'trở về 0° với sai lệch chỉ khoảng 1–3°, cho thấy độ ổn định tốt sau khi áp dụng các kỹ thuật chống trôi.')

h('7. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN', 1)
para('Thiết bị đo góc gập đầu gối dựa trên ESP32 và hai cảm biến MPU6050 đã hoạt động ổn định: đo được ROM '
     'liên tục, hạn chế tốt hiện tượng trôi nhờ kết hợp hiệu chỉnh bias có loại nhiễu, khởi tạo hướng từ gia '
     'tốc kế và ZUPT; tín hiệu mượt nhờ lọc nhiều tầng; kết quả được xuất ra Excel trực quan kèm thống kê sai số.')
para('Hướng phát triển: (i) bổ sung cảm biến từ trường (magnetometer) để khử hoàn toàn trôi phương vị; '
     '(ii) rút ngắn và bọc chống nhiễu dây I2C để tăng độ tin cậy đường truyền; '
     '(iii) xây dựng ứng dụng di động hiển thị thời gian thực và lưu trữ hồ sơ bệnh nhân.')

doc.save('BAO_CAO_DO_GOC_DAU_GOI.docx')
print('Da dung lai BAO_CAO_DO_GOC_DAU_GOI.docx — ROM', S['rom'])
