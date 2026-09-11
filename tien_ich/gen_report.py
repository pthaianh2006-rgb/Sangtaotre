# -*- coding: utf-8 -*-
"""Build a report from measured workbook data without modifying the input."""

from __future__ import annotations

import argparse
from datetime import datetime
from io import BytesIO
import math
import os
from pathlib import Path
import statistics
import tempfile

ROOT = Path(__file__).resolve().parents[1]


def read_samples(source):
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet_name = next((name for name in ("Du lieu", "Dữ liệu") if name in workbook.sheetnames), None)
        if sheet_name is None:
            raise ValueError("File phai co sheet 'Du lieu' hoac 'Dữ liệu' voi cot B la giay, C la goc.")
        samples = []
        for row_number, row in enumerate(workbook[sheet_name].iter_rows(min_row=2, min_col=2, max_col=3, values_only=True), 2):
            elapsed, angle = row
            if elapsed is None and angle is None:
                continue
            try:
                if isinstance(elapsed, bool) or isinstance(angle, bool):
                    raise ValueError
                elapsed, angle = float(elapsed), float(angle)
                if not (math.isfinite(elapsed) and math.isfinite(angle)):
                    raise ValueError
                if elapsed < 0 or not -20 <= angle <= 180:
                    raise ValueError
                if samples and elapsed < samples[-1][0]:
                    raise ValueError
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Du lieu khong hop le tai dong {row_number}.") from exc
            samples.append((elapsed, angle))
        if not samples:
            raise ValueError("File Excel khong co mau do.")
        return samples
    finally:
        workbook.close()


def summarize(samples, target=90.0):
    if not samples:
        raise ValueError("Can it nhat mot mau do.")
    if not math.isfinite(target) or not 0 < target <= 180:
        raise ValueError("Muc tieu phai lon hon 0 va khong qua 180 do.")
    angles = [angle for _, angle in samples]
    duration = samples[-1][0] - samples[0][0]
    count = len(angles)
    sd = statistics.stdev(angles) if count > 1 else None
    return {
        "count": count, "duration": duration,
        "hz": (count - 1) / duration if duration > 0 else None,
        "mean": statistics.fmean(angles), "minimum": min(angles), "rom": max(angles),
        "sd": sd, "sem": sd / math.sqrt(count) if sd is not None else None,
        "target": target, "percent": max(angles) / target * 100,
    }


def build_report(source, output, target=90.0, overwrite=False):
    from docx import Document
    from docx.shared import Inches, Pt
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg

    source, output = Path(source).resolve(), Path(output).resolve()
    if source == output:
        raise ValueError("Tep bao cao phai khac tep du lieu dau vao.")
    if output.suffix.lower() != ".docx":
        raise ValueError("Tep bao cao phai co duoi .docx.")
    if output.exists() and not overwrite:
        raise FileExistsError(f"Tep da ton tai: {output}. Dung --overwrite de thay the.")
    samples = read_samples(source)
    stats = summarize(samples, target)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    document.add_heading("Báo cáo dữ liệu đo góc đầu gối", 0)
    document.add_paragraph(f"Nguồn: {source.name}\nTạo lúc: {datetime.now():%d/%m/%Y %H:%M}")
    document.add_paragraph("Góc gập được giữ nguyên từ file Excel. ROM trong báo cáo là góc gập lớn nhất ghi nhận; mục tiêu là tham số do người dùng lựa chọn.")
    document.add_heading("Thống kê buổi đo", 1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Chỉ số"
    table.rows[0].cells[1].text = "Giá trị"
    fields = (
        ("Số mẫu", "count", ""), ("Thời lượng", "duration", " giây"),
        ("Tốc độ lấy mẫu trung bình", "hz", " Hz"), ("Góc trung bình", "mean", "°"),
        ("Góc nhỏ nhất", "minimum", "°"), ("ROM lớn nhất", "rom", "°"),
        ("Độ lệch chuẩn mẫu", "sd", "°"), ("Sai số chuẩn trung bình", "sem", "°"),
        ("Mục tiêu", "target", "°"), ("ROM / mục tiêu", "percent", "%"),
    )
    for label, key, unit in fields:
        cells = table.add_row().cells
        cells[0].text = label
        value = stats[key]
        cells[1].text = "Không đủ dữ liệu" if value is None else (str(value) if key == "count" else f"{value:.2f}{unit}")

    figure = Figure(figsize=(9, 4), dpi=140)
    FigureCanvasAgg(figure)
    axis = figure.subplots()
    times, angles = zip(*samples)
    axis.plot(times, angles, color="#7c6ae0", linewidth=1.5, marker="." if len(samples) == 1 else None)
    axis.axhline(target, color="#b5573f", linestyle="--", label=f"Mục tiêu {target:g}°")
    axis.set(xlabel="Thời gian (giây)", ylabel="Góc gập (độ)")
    axis.grid(alpha=0.25)
    axis.legend()
    figure.tight_layout()
    chart = BytesIO()
    figure.savefig(chart, format="png")
    chart.seek(0)
    document.add_heading("Góc gập theo thời gian", 1)
    document.add_picture(chart, width=Inches(6.2))
    document.add_paragraph("Độ lệch chuẩn thể hiện sự biến thiên của các mẫu trong buổi đo. Chỉ số này không tự chứng minh độ chính xác của thiết bị; báo cáo không suy ra độ chính xác khi chưa có phép đo tham chiếu.")
    document.add_heading("Dữ liệu đầu buổi (tối đa 20 mẫu)", 1)
    preview = document.add_table(rows=1, cols=3)
    preview.style = "Light Shading Accent 1"
    for cell, heading in zip(preview.rows[0].cells, ("STT", "Thời gian (giây)", "Góc gập (độ)")):
        cell.text = heading
    for index, (elapsed, angle) in enumerate(samples[:20], 1):
        for cell, value in zip(preview.add_row().cells, (str(index), f"{elapsed:.3f}", f"{angle:.1f}")):
            cell.text = value

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, prefix=".report-", suffix=".docx", delete=False) as stream:
            temporary = Path(stream.name)
        document.save(temporary)
        if overwrite:
            os.replace(temporary, output)
        else:
            # Hard-link publication fails atomically if another writer created output.
            os.link(temporary, output)
        return stats
    finally:
        chart.close()
        figure.clear()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Tao bao cao Word tu du lieu Excel KneeROM.")
    parser.add_argument("input", type=Path, help="File .xlsx do GUI/CLI xuat")
    parser.add_argument("--output", type=Path, default=ROOT / "exports" / f"bao_cao_{datetime.now():%Y%m%d_%H%M%S}.docx")
    parser.add_argument("--target", type=float, default=90.0)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args(argv)
    try:
        stats = build_report(args.input, args.output, args.target, args.overwrite)
    except (OSError, ValueError, ImportError) as exc:
        parser.exit(1, f"Khong tao duoc bao cao: {exc}\n")
    print(f"Da tao {args.output.resolve()} ({stats['count']} mau).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
