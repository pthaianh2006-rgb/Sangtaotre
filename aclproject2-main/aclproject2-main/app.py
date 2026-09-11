from flask import g, has_request_context, Flask, render_template, request, Response, jsonify, redirect, url_for, session, send_file
from io import BytesIO
from collections import deque
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
import atexit
import importlib
import logging
from pathlib import Path
import secrets
import sqlite3
import time
import datetime
import math
import os
import json
import socket
import threading
from werkzeug.security import generate_password_hash, check_password_hash

BASE_DIR = Path(__file__).resolve().parent


def _env_bool(name, default=False):
    return os.getenv(name, str(default)).lower() in ("1", "true", "yes", "on")


def _env_int(name, default, minimum, maximum):
    value = int(os.getenv(name, str(default)))
    if not minimum <= value <= maximum:
        raise ValueError(f"{name} must be between {minimum} and {maximum}")
    return value


app = Flask(__name__, template_folder=str(BASE_DIR))
app.config.update(
    SECRET_KEY=os.getenv("ACL_SECRET_KEY") or secrets.token_hex(32),
    DATABASE=str(Path(os.getenv("ACL_DATABASE", str(BASE_DIR / "database.db"))).resolve()),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=_env_bool("ACL_COOKIE_SECURE"),
    MAX_CONTENT_LENGTH=64 * 1024,
    PERMANENT_SESSION_LIFETIME=datetime.timedelta(hours=8),
    HARDWARE_ENABLED=_env_bool("ACL_ENABLE_HARDWARE", True),
    CSRF_ENABLED=True,
)
logger = logging.getLogger(__name__)
_db_init_lock = threading.Lock()
_initialized_databases = set()


def connect_db():
    conn = sqlite3.connect(app.config["DATABASE"], timeout=10)
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA busy_timeout=10000")
    if has_request_context():
        g.setdefault("_db_connections", []).append(conn)
    return conn


@app.teardown_request
def close_request_connections(error=None):
    for conn in g.pop("_db_connections", []):
        conn.close()


def csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return session["csrf_token"]


@app.context_processor
def security_context():
    return {"csrf_token": csrf_token}


@app.before_request
def protect_request():
    if request.endpoint in ("static", "health"):
        return None
    database = app.config["DATABASE"]
    if database not in _initialized_databases:
        with _db_init_lock:
            if database not in _initialized_databases:
                init_db()
                _initialized_databases.add(database)
    if request.method in ("POST", "PUT", "PATCH", "DELETE") and app.config["CSRF_ENABLED"]:
        token = request.headers.get("X-CSRF-Token") or request.form.get("csrf_token", "")
        expected = session.get("csrf_token", "")
        if not expected or not isinstance(token, str) or not secrets.compare_digest(token, expected):
            return jsonify(success=False, message="Phiên làm việc hết hạn. Vui lòng tải lại trang."), 400
    if request.endpoint in _HARDWARE_ENDPOINTS | _PATIENT_ENDPOINTS:
        if "user_id" not in session:
            return jsonify(success=False, message="Vui lòng đăng nhập."), 401
        if session.get("role") != "patient":
            return jsonify(success=False, message="Chỉ tài khoản bệnh nhân được thực hiện thao tác này."), 403
    if request.endpoint in _HARDWARE_ENDPOINTS:
        if not _claim_hardware(session["user_id"]):
            return jsonify(success=False, message="Thiết bị đang được một bệnh nhân khác sử dụng."), 409
        if USE_IMU and request.endpoint in _IMU_ENDPOINTS:
            _start_imu()


@app.after_request
def security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "same-origin"
    if request.endpoint != "static":
        response.headers["Cache-Control"] = "no-store"
    return response

# ===== CAU HINH =====
# Nguon camera cho MediaPipe:
#   0  = webcam laptop
#   1, 2 = webcam ao (vd iPhone qua Iriun/DroidCam)
#   "http://192.168.x.x:8081/video" = stream IP tu app camera tren iPhone
_camera_source = os.getenv("ACL_CAMERA_SOURCE", "1")
CAMERA_SOURCE = int(_camera_source) if _camera_source.isdecimal() else _camera_source
USE_IMU = _env_bool("ACL_USE_IMU", True)
IMU_MODE = os.getenv("ACL_IMU_MODE", "wifi").lower()
if IMU_MODE not in ("wifi", "serial"):
    raise ValueError("ACL_IMU_MODE must be wifi or serial")
IMU_HOST = os.getenv("ACL_IMU_HOST", "192.168.4.1")
IMU_TCP_PORT = _env_int("ACL_IMU_TCP_PORT", 8080, 1, 65535)
IMU_PORT = os.getenv("ACL_IMU_PORT", "COM6")
IMU_BAUD = _env_int("ACL_IMU_BAUD", 115200, 1200, 4000000)
EXPORT_DIR = os.getenv("ACL_EXPORT_DIR", str(BASE_DIR / "exports"))

def _save_and_send(wb, filename):
    """Serialize once and optionally keep an atomic local export copy."""
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
    bio = BytesIO()
    wb.save(bio)
    if EXPORT_DIR:
        temporary = None
        try:
            directory = Path(EXPORT_DIR)
            directory.mkdir(parents=True, exist_ok=True)
            temporary = directory / f".{secrets.token_hex(8)}.tmp"
            temporary.write_bytes(bio.getvalue())
            temporary.replace(directory / Path(filename).name)
        except OSError:
            logger.warning("Could not save local export", exc_info=True)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def init_db():
    Path(app.config["DATABASE"]).parent.mkdir(parents=True, exist_ok=True)
    conn = connect_db()
    cursor = conn.cursor()
    
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        role TEXT NOT NULL, -- 'doctor' hoặc 'patient'
        name TEXT NOT NULL,
        age INTEGER,
        week INTEGER,
        doctor_advice TEXT
    )
    """)
   
   
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS measurements (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        patient_id INTEGER,
        angle REAL,
        time TEXT,
        source TEXT DEFAULT 'camera',
        FOREIGN KEY(patient_id) REFERENCES users(id)
    )
    """)
    measurement_columns = {row[1] for row in cursor.execute("PRAGMA table_info(measurements)")}
    if "source" not in measurement_columns:
        cursor.execute("ALTER TABLE measurements ADD COLUMN source TEXT DEFAULT 'camera'")

    # Migration: muc tieu ROM tuy chinh cho tung benh nhan (vd 90 do trong 14 ngay)
    user_columns = {row[1] for row in cursor.execute("PRAGMA table_info(users)")}
    for col, ddl in (("rom_goal", "INTEGER"), ("rom_goal_days", "INTEGER"), ("rom_goal_start", "TEXT")):
        if col not in user_columns:
            cursor.execute("ALTER TABLE users ADD COLUMN %s %s" % (col, ddl))
    
    
    admin_password = os.getenv("ACL_ADMIN_PASSWORD")
    admin_username = os.getenv("ACL_ADMIN_USERNAME", "admin").strip()
    if admin_password and not cursor.execute("SELECT 1 FROM users WHERE username=?", (admin_username,)).fetchone():
        if len(admin_password) < 10:
            conn.close()
            raise ValueError("ACL_ADMIN_PASSWORD must contain at least 10 characters")
        hashed_pw = generate_password_hash(admin_password)
        cursor.execute("""
            INSERT INTO users (username, password, role, name, age, week, doctor_advice) 
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (admin_username, hashed_pw, "doctor", "Bác sĩ Quản trị", 40, 0, ""))
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_measurements_patient_id_id ON measurements(patient_id, id DESC)")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_measurements_patient_source_time ON measurements(patient_id, source, time)")
    conn.commit()
    conn.close()

camera = None
pose = None
cv2 = None
mp_pose = None
POSE_W = 640                                   # do rong ANH DUA VAO MediaPipe (lon hon -> ro khop hon)
VIS_TH = 0.3                                    # nguong visibility (thap -> it rot khop)
HOLD_S = 0.5                                    # giu landmark cu bao lau khi rot tam thoi (giay)
MAX_JUMP = 0.18                                 # dau goi duoc phep nhay toi da/frame (chuan hoa) - chong nhay nguoi/nhieu
RELOCK_S = 0.6                                  # mat dau goi qua lau -> cho phep KHOA LAI o vi tri moi
# Quay NGHIENG chi do 1 chan (chan ro nhat) -> dung 1 khoa "active"
_leg_persist = {"active": None}                 # luu landmark gan nhat de "giu" khi rot khop
_leg_lastpos = {"active": None}                 # vi tri dau goi gan nhat (x,y,t) de loc nhay theo chuyen dong
_leg_flexbuf = {"active": deque(maxlen=5)}      # bo dem goc de loc trung vi (diet gai nhon)
# NEO HONG: ngoi tua ghe -> hong gan nhu DUNG YEN nhung hay bi lung ghe che/danh lua.
# Giu vi tri hong on dinh (EMA cham); khi hong bi che/nhay xa -> dung lai hong da neo.
_hip_anchor = {"active": None}                  # vi tri hong da neo (hx, hy) chuan hoa
HIP_EMA = 0.12                                   # toc do cap nhat hong (thap = on dinh, vi hong it di chuyen)
HIP_JUMP = 0.12                                  # hong nhay xa hon nguong nay -> coi nhu bi ghe danh lua, GIU hong cu

# --- Goc do tu CAMERA ---
current_angle = 0
max_rom_session = 0
cam_leg = "auto"         # chan do bang camera: "left" | "right" | "auto" (chi de dat ten chan)
crowd_mode = False       # che do DONG NGUOI: chi xu ly nguoi trong VUNG GIUA (bo qua nguoi xung quanh)
ROI_L, ROI_R, ROI_T, ROI_B = 0.22, 0.78, 0.0, 1.0   # vung giua khi bat che do dong nguoi (~56% rong)

class OneEuroFilter:
    """One-Euro filter - chuan vang loc tin hieu chuyen dong nguoi:
       dung yen thi CUC MUOT (giat ~0), di chuyen thi BAM NHANH (it tre).
       Tot hon Kalman cho viec nay vi Kalman bi vot lo (overshoot) khi dung dot ngot."""
    def __init__(self, mincutoff=0.3, beta=0.02, dcutoff=1.0):
        self.mincutoff = mincutoff       # nho hon = muot hon luc dung yen
        self.beta = beta                 # lon hon = bam nhanh hon luc di chuyen (nhung giat hon)
        self.dcutoff = dcutoff
        self.x_prev = None; self.dx_prev = 0.0; self.t_prev = None
    def _alpha(self, cutoff, dt):
        tau = 1.0 / (2.0 * math.pi * cutoff)
        return 1.0 / (1.0 + tau / dt)
    def update(self, x, t):
        if self.x_prev is None:
            self.x_prev = x; self.t_prev = t; return x
        dt = t - self.t_prev; self.t_prev = t
        if dt <= 0 or dt > 0.5: dt = 0.033
        dx = (x - self.x_prev) / dt
        a_d = self._alpha(self.dcutoff, dt)
        dx_hat = a_d * dx + (1.0 - a_d) * self.dx_prev
        self.dx_prev = dx_hat
        cutoff = self.mincutoff + self.beta * abs(dx_hat)   # cutoff thich nghi theo toc do
        a = self._alpha(cutoff, dt)
        x_hat = a * x + (1.0 - a) * self.x_prev
        self.x_prev = x_hat
        return x_hat

cam_filter_L = OneEuroFilter(mincutoff=0.15, beta=0.02)  # bo loc chan TRAI (mincutoff thap -> rat muot luc yen)
cam_filter_R = OneEuroFilter(mincutoff=0.15, beta=0.02)  # bo loc chan PHAI

# --- Goc do tu CAM BIEN IMU (rieng biet de chay song song + so sanh) ---
imu_angle = 0
imu_max_rom = 0
imu_jitter = 0.0                  # do dao dong (rung) cua tin hieu gan day
_imu_buf = deque(maxlen=15)       # bo dem goc gan nhat de tinh jitter
imu_calibrating = False           # dang lay gyro bias luc boot (giu yen)
imu_last_data = 0.0               # thoi diem nhan duoc dong du lieu gan nhat
current_port = IMU_PORT            # cong dang dung (USB hoac Bluetooth) - doi duoc luc chay
_want_reconnect = False            # co hieu: yeu cau ket noi lai cong khac
imu_serial = None
imu_connected = False

_state_lock = threading.RLock()
_worker_lock = threading.Lock()
_imu_write_lock = threading.Lock()
_frame_condition = threading.Condition()
_stop_event = threading.Event()
_camera_thread = None
_imu_thread = None
_latest_frame = None
_frame_number = 0
_camera_error = None
camera_last_data = 0.0
_hardware_owner = None
_hardware_last_seen = 0.0
HARDWARE_LEASE_SECONDS = 15.0
_HARDWARE_ENDPOINTS = {
    "video_feed", "get_angle", "get_max_rom", "set_leg", "toggle_crowd",
    "get_imu_angle", "get_imu_max_rom", "imu_status", "imu_reconnect",
    "set_imu_mode", "imu_ports", "imu_connect", "imu_cmd", "reset_rom", "save_rom", "game_save",
}
_IMU_ENDPOINTS = {"get_imu_angle", "get_imu_max_rom", "imu_status", "imu_reconnect", "set_imu_mode", "imu_connect", "imu_cmd"}
_PATIENT_ENDPOINTS = {"set_goal", "game_save", "save_rom", "get_measurements", "delete_all_me"}


def _reset_tracking():
    global current_angle, max_rom_session, imu_angle, imu_max_rom, camera_last_data, imu_last_data
    global cam_filter_L, cam_filter_R, _latest_frame, imu_jitter
    current_angle = max_rom_session = imu_angle = imu_max_rom = 0
    camera_last_data = imu_last_data = 0.0
    imu_jitter = 0.0
    _imu_buf.clear()
    _leg_persist["active"] = _leg_lastpos["active"] = _hip_anchor["active"] = None
    _leg_flexbuf["active"].clear()
    cam_filter_L = OneEuroFilter(mincutoff=0.15, beta=0.02)
    cam_filter_R = OneEuroFilter(mincutoff=0.15, beta=0.02)
    _latest_frame = None


def _claim_hardware(patient_id):
    global _hardware_owner, _hardware_last_seen
    now = time.monotonic()
    with _state_lock:
        if _hardware_owner != patient_id:
            if _hardware_owner is not None and now - _hardware_last_seen < HARDWARE_LEASE_SECONDS:
                return False
            _reset_tracking()
            _hardware_owner = patient_id
        _hardware_last_seen = now
    return True


def _hardware_active():
    return _hardware_owner is not None and time.monotonic() - _hardware_last_seen < HARDWARE_LEASE_SECONDS


def _start_imu():
    global _imu_thread
    if not USE_IMU or not app.config["HARDWARE_ENABLED"]:
        return
    with _worker_lock:
        if _imu_thread is None or not _imu_thread.is_alive():
            _imu_thread = threading.Thread(target=imu_reader, name="acl-imu", daemon=True)
            _imu_thread.start()


def shutdown_resources():
    """Stop workers and release handles without opening hardware during import."""
    _stop_event.set()
    with _frame_condition:
        _frame_condition.notify_all()
    with _imu_write_lock:
        if imu_serial is not None:
            try:
                imu_serial.close()
            except OSError:
                pass
    for worker in (_camera_thread, _imu_thread):
        if worker and worker is not threading.current_thread():
            worker.join(timeout=6)


atexit.register(shutdown_resources)

def _process_imu_line(line):
    """Xu ly 1 dong tu ESP32: nhan biet boot/bias + tinh goc gap. Dung chung WiFi & Serial."""
    global imu_angle, imu_max_rom, imu_jitter, imu_calibrating, imu_last_data
    if not line:
        return
    if line.startswith("[INFO]") or line.startswith("[WARN]"):
        low = line.lower()
        if "giu yen" in low or "rung" in low:
            imu_calibrating = True
        elif "san sang" in low:
            imu_calibrating = False
        return
    knee = None
    if line.startswith("Knee:"):                     # "Knee:83.4 [axis]"
        try: knee = float(line.split(":")[1].split()[0])
        except (ValueError, IndexError): pass
    elif "," in line:                                 # che do log "millis,knee"
        parts = line.split(",")
        if len(parts) == 2:
            try: knee = float(parts[1])
            except ValueError: pass
    if knee is not None and math.isfinite(knee) and 0 <= knee <= 200:
        with _state_lock:
            imu_last_data = time.monotonic()
            imu_calibrating = False
            imu_angle = round(180 - knee, 1)
            if _hardware_active():
                imu_max_rom = max(imu_max_rom, imu_angle)
            _imu_buf.append(imu_angle)
            if len(_imu_buf) >= 5:
                imu_jitter = max(_imu_buf) - min(_imu_buf)

def imu_reader():
    """Bounded reads and deterministic cleanup for serial and TCP reconnects."""
    global imu_serial, imu_connected, imu_calibrating, _want_reconnect
    while not _stop_event.is_set():
        if not _hardware_active():
            _stop_event.wait(0.5)
            continue
        connection = None
        try:
            mode, opened = IMU_MODE, current_port
            _want_reconnect = False
            if mode == "wifi":
                connection = socket.create_connection((IMU_HOST, IMU_TCP_PORT), timeout=3)
                connection.settimeout(1)
            else:
                serial_module = importlib.import_module("serial")
                connection = serial_module.Serial(opened, IMU_BAUD, timeout=1, write_timeout=1)
            with _imu_write_lock:
                imu_serial = connection
                imu_connected = True
                imu_calibrating = True
            pending = b""
            discarding = False
            while not _stop_event.is_set() and _hardware_active():
                if _want_reconnect or mode != IMU_MODE or (mode == "serial" and opened != current_port):
                    break
                try:
                    raw = connection.recv(4096) if mode == "wifi" else connection.read(256)
                except socket.timeout:
                    continue
                if not raw and mode == "wifi":
                    raise OSError("ESP32 disconnected")
                pending += raw
                while b"\n" in pending:
                    line, pending = pending.split(b"\n", 1)
                    if not discarding and len(line) <= 1024:
                        _process_imu_line(line.decode("utf-8", errors="ignore").strip())
                    discarding = False
                if len(pending) > 1024:
                    pending = b""
                    discarding = True
        except (OSError, ImportError, ValueError):
            logger.warning("IMU connection failed; retrying in 3 seconds", exc_info=True)
        finally:
            with _imu_write_lock:
                imu_connected = False
                imu_serial = None
                if connection is not None:
                    try:
                        connection.close()
                    except OSError:
                        pass
        _stop_event.wait(3)

def imu_send(cmd):
    """Gui lenh dieu khien (C/F/L/R) toi ESP32 (WiFi hoac Serial)."""
    global imu_serial, imu_connected
    with _imu_write_lock:
        if imu_serial and imu_connected:
            try:
                data = cmd.encode("ascii")
                if hasattr(imu_serial, "sendall"):
                    imu_serial.sendall(data)
                else:
                    imu_serial.write(data)
                return True
            except OSError:
                imu_connected = False
    return False

def calculate_angle(a, b, c):
    if math.hypot(a[0] - b[0], a[1] - b[1]) < 1e-8 or math.hypot(c[0] - b[0], c[1] - b[1]) < 1e-8:
        raise ValueError("Landmarks overlap")
    radians = math.atan2(c[1] - b[1], c[0] - b[0]) - math.atan2(a[1] - b[1], a[0] - b[0])
    angle = abs(math.degrees(radians))
    if angle > 180:
        angle = 360 - angle
    return angle

# --- ĐĂNG NHẬP / ĐĂNG KÝ ---

@app.route("/")
def home():
    if "user_id" in session:
        if session["role"] == "doctor":
            return redirect(url_for("admin"))
        return redirect(url_for("dashboard"))
    return render_template("home.html")

@app.route("/register", methods=["POST"])
def register():
    try:
        username = request.form["username"].strip()
        password = request.form["password"]
        name = request.form["name"].strip()
        age = int(request.form["age"])
        week = int(request.form["week"])
        if not 3 <= len(username) <= 64 or any(ch.isspace() for ch in username):
            raise ValueError("Tên đăng nhập phải có 3–64 ký tự, không chứa khoảng trắng.")
        if not 10 <= len(password) <= 256:
            raise ValueError("Mật khẩu phải có 10–256 ký tự.")
        if not 1 <= len(name) <= 120 or not 1 <= age <= 120 or not 0 <= week <= 104:
            raise ValueError("Họ tên, tuổi hoặc tuần phục hồi không hợp lệ.")
        
        hashed_password = generate_password_hash(password)
        
        conn = connect_db()
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO users (username, password, role, name, age, week, doctor_advice)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (username, hashed_password, "patient", name, age, week, ""))
        conn.commit()
        msg, category = "Đăng ký thành công! Vui lòng đăng nhập.", "success"
    except sqlite3.IntegrityError:
        msg, category = "Tên đăng nhập đã tồn tại!", "danger"
    except (ValueError, KeyError):
        return render_template("home.html", msg="Thông tin không hợp lệ. Kiểm tra tên đăng nhập, mật khẩu (ít nhất 10 ký tự), tuổi và tuần phục hồi.", category="danger"), 400
    except sqlite3.Error:
        logger.exception("Registration failed")
        return render_template("home.html", msg="Chưa thể tạo tài khoản. Vui lòng thử lại.", category="danger"), 503
    finally:
        if 'conn' in locals(): conn.close()
        
    return render_template("home.html", msg=msg, category=category)

@app.route("/login", methods=["POST"])
def login():
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    if not username or len(username) > 64 or not password or len(password) > 256:
        return render_template("home.html", msg="Sai tài khoản hoặc mật khẩu!", category="danger"), 400
    
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM users WHERE username = ?", (username,))
    user = cursor.fetchone()
    conn.close()
    
    if user and check_password_hash(user[2], password):
        session.clear()
        session.permanent = True
        session["user_id"] = user[0]
        session["username"] = user[1]
        session["role"] = user[3]
        session["name"] = user[4]
        
        if user[3] == "doctor":
            return redirect(url_for("admin"))
        return redirect(url_for("dashboard"))
    else:
        return render_template("home.html", msg="Sai tài khoản hoặc mật khẩu!", category="danger")

@app.route("/logout", methods=["GET", "POST"])
def logout():
    global _hardware_owner, _hardware_last_seen
    with _state_lock:
        if _hardware_owner == session.get("user_id"):
            _hardware_owner = None
            _hardware_last_seen = 0.0
            _reset_tracking()
    session.clear()
    return redirect(url_for("home"))



@app.route("/dashboard")
def dashboard():
    if "user_id" not in session or session["role"] != "patient":
        return redirect(url_for("home"))
        
    conn = connect_db()
    cursor = conn.cursor()
   
    cursor.execute("SELECT age, week, doctor_advice FROM users WHERE id = ?", (session["user_id"],))
    user_info = cursor.fetchone()
    
    
    cursor.execute("SELECT angle, time, source, id FROM measurements WHERE patient_id = ? ORDER BY id DESC", (session["user_id"],))
    history_rows = cursor.fetchall()
    conn.close()

    if user_info is None:
        session.clear()
        return redirect(url_for("home"))
    age, week, custom_advice = user_info
    week = week or 0

    if week <= 2:
        phase, target = "Giai đoạn 1", "0° - 90°"
        recommendation = {"muc_tieu": "Kiểm soát đau và sưng", "loi_khuyen": ["Chườm lạnh sau tập", "Co cơ tĩnh cơ tứ đầu", "Không gập quá 90°"]}
    elif week <= 6:
        phase, target = "Giai đoạn 2", "90° - 125°"
        recommendation = {"muc_tieu": "Tăng ROM và sức cơ", "loi_khuyen": ["Đạp xe tại chỗ", "Tập squat nông", "Tăng biên độ từ từ"]}
    else:
        phase, target = "Giai đoạn 3", "125° - 135°"
        recommendation = {"muc_tieu": "Khôi phục vận động toàn diện", "loi_khuyen": ["Đi bộ nhanh", "Chạy bộ nhẹ", "Tập thăng bằng chân nâng cao"]}

  
    if custom_advice:
        recommendation["loi_khuyen"].insert(0, f"⭐ LỜI KHUYÊN BÁC SĨ: {custom_advice}")

    return render_template("dashboard.html", name=session["name"], age=age, week=week, phase=phase, target=target, recommendation=recommendation, history=history_rows, use_imu=USE_IMU, imu_mode=IMU_MODE)

@app.route("/profile")
def profile():
    if "user_id" not in session or session["role"] != "patient":
        return redirect(url_for("home"))
        
    pid = session["user_id"]
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id, name, age, week, username FROM users WHERE id=?", (pid,))
    user = cursor.fetchone()
    
    if not user:
        conn.close()
        return "Không tìm thấy hồ sơ!", 404
        
    # Lay ROM lon nhat moi ngay, tach rieng Camera va Cam bien de ve 2 duong
    cursor.execute("""
        SELECT SUBSTR(time,1,10) AS d,
               MAX(CASE WHEN source='camera' THEN angle END),
               MAX(CASE WHEN source='imu'    THEN angle END)
        FROM measurements WHERE patient_id=? GROUP BY d ORDER BY d ASC
    """, (pid,))
    chart_data = cursor.fetchall()
    conn.close()
    dates = [row[0] if row[0] else "" for row in chart_data]
    cam_angles = [row[1] for row in chart_data]
    imu_angles = [row[2] for row in chart_data]

    return render_template("profile.html", user=user,
                           dates=json.dumps(dates, ensure_ascii=False),
                           cam_angles=json.dumps(cam_angles),
                           imu_angles=json.dumps(imu_angles))

@app.route("/my_records")
def my_records():
    """Trang ho so: benh nhan xem lai qua trinh tap theo tung ngay."""
    if "user_id" not in session or session["role"] != "patient":
        return redirect(url_for("home"))
    pid = session["user_id"]
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT name, age, week, doctor_advice FROM users WHERE id=?", (pid,))
    info = cursor.fetchone()
    # Tong hop theo NGAY: so luot, ROM camera tot nhat, ROM cam bien tot nhat
    cursor.execute("""
        SELECT SUBSTR(time,1,10) AS d, COUNT(*),
               MAX(CASE WHEN source='camera' THEN angle END),
               MAX(CASE WHEN source='imu'    THEN angle END)
        FROM measurements WHERE patient_id=? GROUP BY d ORDER BY d DESC
    """, (pid,))
    days = cursor.fetchall()
    # Chi tiet tung luot do
    cursor.execute("SELECT time, angle, source FROM measurements WHERE patient_id=? ORDER BY id DESC", (pid,))
    allm = cursor.fetchall()
    conn.close()
    return render_template("my_records.html", info=info, days=days, allm=allm)



def _process_camera_frame(frame):
    global current_angle, max_rom_session, camera_last_data
    h, w, _ = frame.shape

    # CHE DO DONG NGUOI: chi cat VUNG GIUA dua vao MediaPipe -> nguoi xung quanh "vo hinh"
    if crowd_mode:
        rx0, rx1 = int(ROI_L * w), int(ROI_R * w)
        ry0, ry1 = int(ROI_T * h), int(ROI_B * h)
    else:
        rx0, ry0, rx1, ry1 = 0, 0, w, h
    rw, rh = max(1, rx1 - rx0), max(1, ry1 - ry0)
    roi = frame[ry0:ry1, rx0:rx1]

    # XU LY anh (chua lat) -> MediaPipe phan biet trai/phai giai phau chinh xac
    small = cv2.resize(roi, (POSE_W, max(1, int(rh * POSE_W / rw))))
    rgb = cv2.cvtColor(small, cv2.COLOR_BGR2RGB)
    rgb.flags.writeable = False           # bo qua copy -> nhanh hon
    result = pose.process(rgb)
    # Lat anh de HIEN THI nhu guong (cho tu nhien khi tap)
    frame = cv2.flip(frame, 1)
    if crowd_mode:                        # ve khung VUNG DO (toa do da lat: x -> w-x)
        cv2.rectangle(frame, (w - rx1, ry0), (w - rx0, ry1), (255, 180, 0), 2)
        cv2.putText(frame, "VUNG DO - dung trong khung", (w - rx1 + 6, ry0 + 26),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.6, (255, 180, 0), 2)

    if result.pose_landmarks:
        lm = result.pose_landmarks.landmark
        P = mp_pose.PoseLandmark
        tnow = time.monotonic()
        # === QUAY NGHIENG: KHONG dung nhan Trai/Phai cua MediaPipe (rat hay nham vi
        # 2 chan chong nhau). Luon do CHAN RO NHAT trong khung = chan gan camera =
        # chan dang quay. Nut Trai/Phai chi de DAT TEN chan luc ghi ho so. ===
        cand = []
        for pset in ((P.LEFT_HIP, P.LEFT_KNEE, P.LEFT_ANKLE),
                     (P.RIGHT_HIP, P.RIGHT_KNEE, P.RIGHT_ANKLE)):
            Hh, Kk, Aa = lm[pset[0].value], lm[pset[1].value], lm[pset[2].value]
            # GATING theo GOI + CO CHAN (phan di chuyen, o duoi ghe nen luon ro);
            # KHONG xet hong vi ngoi tua ghe hay bi lung ghe che.
            vis = min(Kk.visibility, Aa.visibility)
            cand.append((vis, Hh, Kk, Aa))
        vis, Hh, Kk, Aa = max(cand, key=lambda c: c[0])   # chan ro nhat (visibility cao nhat)

        # ten + mau theo lua chon chan (chi de ghi nhan, khong anh huong nhan dien)
        if cam_leg == "right":
            name, color, filt = "Phai", (0, 200, 255), cam_filter_R
        elif cam_leg == "left":
            name, color, filt = "Trai", (0, 255, 0), cam_filter_L
        else:
            name, color, filt = "Chan do", (0, 255, 0), cam_filter_L

        key = "active"
        solid = False; draw = False
        if vis > VIS_TH:
            hx = (rx0 + Hh.x * rw) / w; hy = (ry0 + Hh.y * rh) / h
            kx = (rx0 + Kk.x * rw) / w; ky = (ry0 + Kk.y * rh) / h
            ax = (rx0 + Aa.x * rw) / w; ay = (ry0 + Aa.y * rh) / h
            # --- NEO HONG: chong lung ghe che/danh lua diem hong ---
            anc = _hip_anchor.get(key)
            if anc is None:
                _hip_anchor[key] = [hx, hy]            # khoi tao neo
            else:
                hip_ok = (Hh.visibility > VIS_TH and
                          ((hx - anc[0]) ** 2 + (hy - anc[1]) ** 2) ** 0.5 <= HIP_JUMP)
                if hip_ok:                             # hong dang ro & gan -> cap nhat cham (EMA)
                    anc[0] += (hx - anc[0]) * HIP_EMA
                    anc[1] += (hy - anc[1]) * HIP_EMA
                # else: hong bi che/nhay xa -> GIU nguyen neo (hong ngoi it di chuyen)
            hx, hy = _hip_anchor[key]                  # dung hong da neo cho tinh goc & ve
            # LOC THEO CHUYEN DONG: chi nhan neu dau goi khong nhay qua xa frame truoc
            lp = _leg_lastpos.get(key)
            if lp is None or (tnow - lp[2]) >= RELOCK_S or \
               ((kx - lp[0]) ** 2 + (ky - lp[1]) ** 2) ** 0.5 <= MAX_JUMP:
                _leg_lastpos[key] = (kx, ky, tnow)
                raw_angle = calculate_angle([hx * w, hy * h], [kx * w, ky * h], [ax * w, ay * h])
                flex = max(0, min(170, 180 - raw_angle))   # DUOI THANG = 0, gap cang nhieu cang TANG
                buf = _leg_flexbuf[key]; buf.append(flex)
                med = sorted(buf)[len(buf) // 2]      # MEDIAN-of-5: diet gai nhon
                sm = filt.update(med, tnow)           # One-Euro lam muot
                _leg_persist[key] = {
                    "hip":   (int((1 - hx) * w), int(hy * h)),
                    "knee":  (int((1 - kx) * w), int(ky * h)),
                    "ankle": (int((1 - ax) * w), int(ay * h)),
                    "sm": sm, "t": tnow}
                solid = True; draw = True
        if not draw:
            # rot khop / nhay xa -> GIU landmark cu neu con moi (< HOLD_S)
            p = _leg_persist.get(key)
            if p and (tnow - p["t"]) <= HOLD_S:
                sm = p["sm"]; draw = True
        if draw:
            p = _leg_persist[key]
            th = 6 if solid else 3
            cv2.line(frame, p["hip"], p["knee"], color, th)
            cv2.line(frame, p["knee"], p["ankle"], color, th)
            cv2.circle(frame, p["knee"], 9, (0, 0, 255), -1)
            cv2.putText(frame, f"{name}: {int(round(sm))} do", (p["knee"][0]-60, p["knee"][1]-15),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.8, color, 3)
            current_angle = int(round(sm))
            if solid: camera_last_data = time.monotonic()
            if current_angle > max_rom_session: max_rom_session = current_angle

    ret, buffer = cv2.imencode('.jpg', frame, [cv2.IMWRITE_JPEG_QUALITY, 80])
    return buffer.tobytes() if ret else None

def _open_camera():
    global cv2, mp_pose, camera, pose
    cv2 = importlib.import_module("cv2")
    mp = importlib.import_module("mediapipe")
    if not hasattr(mp, "solutions"):
        raise RuntimeError("MediaPipe solutions API unavailable; install the supported camera requirements")
    mp_pose = mp.solutions.pose
    camera = cv2.VideoCapture(CAMERA_SOURCE)
    if not camera.isOpened():
        raise OSError("Camera unavailable")
    camera.set(cv2.CAP_PROP_FOURCC, cv2.VideoWriter_fourcc(*"MJPG"))
    camera.set(cv2.CAP_PROP_FRAME_WIDTH, 1280)
    camera.set(cv2.CAP_PROP_FRAME_HEIGHT, 720)
    camera.set(cv2.CAP_PROP_FPS, 30)
    camera.set(cv2.CAP_PROP_BUFFERSIZE, 1)
    pose = mp_pose.Pose(model_complexity=1, smooth_landmarks=True,
                        min_detection_confidence=0.5, min_tracking_confidence=0.5)


def _camera_worker():
    global _latest_frame, _frame_number, _camera_error, camera, pose
    try:
        _open_camera()
        failures = 0
        while not _stop_event.is_set() and _hardware_active():
            started = time.monotonic()
            success, frame = camera.read()
            if not success:
                failures += 1
                if failures >= 10:
                    raise OSError("Camera stopped returning frames")
                _stop_event.wait(0.1)
                continue
            failures = 0
            with _state_lock:
                jpeg = _process_camera_frame(frame)
            if jpeg is not None:
                with _frame_condition:
                    _latest_frame = jpeg
                    _frame_number += 1
                    _camera_error = None
                    _frame_condition.notify_all()
            _stop_event.wait(max(0, 1 / 30 - (time.monotonic() - started)))
    except Exception:
        logger.exception("Camera worker stopped")
        _camera_error = "Không thể mở hoặc xử lý camera. Kiểm tra thiết bị và cấu hình."
    finally:
        if camera is not None:
            camera.release()
            camera = None
        if pose is not None:
            pose.close()
            pose = None
        with _frame_condition:
            _frame_condition.notify_all()


def _start_camera():
    global _camera_thread, _camera_error
    if not app.config["HARDWARE_ENABLED"] or _stop_event.is_set():
        return False
    with _worker_lock:
        if _camera_thread is None or not _camera_thread.is_alive():
            _camera_error = None
            _camera_thread = threading.Thread(target=_camera_worker, name="acl-camera", daemon=True)
            _camera_thread.start()
    return True


def generate_frames(patient_id):
    """Fan out the latest JPEG: one capture/inference worker for all viewers."""
    previous = -1
    while not _stop_event.is_set():
        with _state_lock:
            if _hardware_owner != patient_id or not _hardware_active():
                return
        with _frame_condition:
            _frame_condition.wait_for(
                lambda: (_latest_frame is not None and _frame_number != previous) or _stop_event.is_set() or _camera_error,
                timeout=2,
            )
            if _stop_event.is_set() or _camera_error:
                return
            if _frame_number == previous or _latest_frame is None:
                continue
            previous, jpeg = _frame_number, _latest_frame
        with _state_lock:
            if _hardware_owner != patient_id:
                return
        yield b"--frame\r\nContent-Type: image/jpeg\r\n\r\n" + jpeg + b"\r\n"


@app.route("/health")
def health():
    return jsonify(status="ok", app="KneeROM")


@app.route('/video_feed')
def video_feed():
    if "user_id" not in session: return "Unauthorized", 401
    if not _start_camera():
        return jsonify(success=False, message="Camera đang bị tắt trong cấu hình."), 503
    with _frame_condition:
        _frame_condition.wait_for(lambda: _latest_frame is not None or _camera_error is not None, timeout=8)
    if _latest_frame is None:
        return jsonify(success=False, message=_camera_error or "Camera chưa sẵn sàng."), 503
    return Response(generate_frames(session["user_id"]), mimetype='multipart/x-mixed-replace; boundary=frame')

@app.route("/get_angle")
def get_angle():
    _start_camera()
    with _state_lock:
        age = time.monotonic() - camera_last_data if camera_last_data else None
        return jsonify(angle=current_angle, tracking=age is not None and age < 1, age=age,
                       camera_available=_camera_error is None, message=_camera_error)

@app.route("/get_max_rom")
def get_max_rom():
    with _state_lock:
        return jsonify(max_rom=max_rom_session)

@app.route("/set_leg/<leg>", methods=["POST"])
def set_leg(leg):
    """Chon chan do bang camera: left / right / auto."""
    global cam_leg
    if "user_id" not in session: return jsonify({"success": False}), 401
    if leg in ("left", "right", "auto"):
        with _state_lock:
            cam_leg = leg
            _leg_persist["active"] = _leg_lastpos["active"] = _hip_anchor["active"] = None
            _leg_flexbuf["active"].clear()
        return jsonify({"success": True, "leg": leg})
    return jsonify({"success": False, "message": "Chan khong hop le"}), 400


@app.route("/toggle_crowd", methods=["POST"])
def toggle_crowd():
    """Bat/tat che do dong nguoi (chi do nguoi dung trong vung giua)."""
    global crowd_mode
    if "user_id" not in session: return jsonify({"success": False}), 401
    with _state_lock:
        crowd_mode = not crowd_mode
        _leg_persist["active"] = _leg_lastpos["active"] = _hip_anchor["active"] = None
        _leg_flexbuf["active"].clear()
    return jsonify({"success": True, "crowd": crowd_mode})

@app.route("/get_imu_angle")
def get_imu_angle():
    with _state_lock:
        age = time.monotonic() - imu_last_data if imu_last_data else None
        return jsonify(angle=imu_angle, receiving=imu_connected and age is not None and age < 3, age=age)

@app.route("/get_imu_max_rom")
def get_imu_max_rom():
    with _state_lock:
        return jsonify(max_rom=imu_max_rom)

@app.route("/imu_status")
def imu_status():
    receiving = imu_connected and imu_last_data > 0 and (time.monotonic() - imu_last_data < 3.0)
    where = f"{IMU_HOST}:{IMU_TCP_PORT}" if IMU_MODE == "wifi" else current_port
    return jsonify({"connected": imu_connected, "receiving": receiving, "use_imu": USE_IMU,
                    "jitter": round(imu_jitter, 1), "calibrating": imu_calibrating,
                    "mode": IMU_MODE, "where": where, "port": current_port})

@app.route("/imu_reconnect", methods=["POST"])
def imu_reconnect():
    """Buoc ket noi lai ESP32 (dung cho ca WiFi va Serial)."""
    global _want_reconnect
    if "user_id" not in session: return jsonify({"success": False}), 401
    _want_reconnect = True
    return jsonify({"success": True})

@app.route("/set_imu_mode/<mode>", methods=["POST"])
def set_imu_mode(mode):
    """Chuyen kieu ket noi IMU luc dang chay: 'serial' (USB) hoac 'wifi'."""
    global IMU_MODE, _want_reconnect
    if "user_id" not in session: return jsonify({"success": False}), 401
    if mode not in ("serial", "wifi"):
        return jsonify({"success": False, "message": "Kieu khong hop le"})
    IMU_MODE = mode
    _want_reconnect = True               # ngat ket noi cu -> luong tu noi lai theo mode moi
    return jsonify({"success": True, "mode": mode})

@app.route("/imu_ports")
def imu_ports():
    """Liet ke cac cong COM (USB + Bluetooth) de nguoi dung chon."""
    from serial.tools import list_ports
    out = []
    for p in list_ports.comports():
        desc = p.description or ""
        out.append({"device": p.device, "desc": desc, "bluetooth": "bluetooth" in desc.lower()})
    return jsonify({"ports": out, "current": current_port})

@app.route("/imu_connect/<port>", methods=["POST"])
def imu_connect(port):
    """Doi cong ket noi (USB <-> Bluetooth) ngay luc dang chay."""
    global current_port, _want_reconnect
    if "user_id" not in session: return jsonify({"success": False, "message": "Chưa đăng nhập"}), 401
    current_port = port
    _want_reconnect = True
    return jsonify({"success": True, "port": port})

@app.route("/imu_cmd/<cmd>", methods=["POST"])
def imu_cmd(cmd):
    if "user_id" not in session: return jsonify({"success": False, "message": "Chua dang nhap"}), 401
    cmd = cmd.upper()
    if cmd not in ("C", "L"):
        return jsonify({"success": False, "message": "Lenh khong hop le"})
    ok = imu_send(cmd)
    return jsonify({"success": ok, "message": ("Da gui lenh " + cmd) if ok else "Cam bien IMU chua ket noi"})

@app.route("/get_measurements")
def get_measurements():
    if "user_id" not in session: return jsonify([])
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT angle, time FROM measurements WHERE patient_id = ? ORDER BY id DESC", (session["user_id"],))
    rows = cursor.fetchall()
    conn.close()
    return jsonify([{"angle": r[0], "time": r[1]} for r in rows])

@app.route("/reset_rom/<source>", methods=["POST"])
def reset_rom(source):
    global max_rom_session, imu_max_rom, camera_last_data, imu_last_data
    with _state_lock:
        if source == "camera":
            max_rom_session = 0
            camera_last_data = 0.0
        elif source == "imu":
            imu_max_rom = 0
            imu_last_data = 0.0
        else:
            return jsonify(success=False, message="Nguồn không hợp lệ."), 400
    return jsonify(success=True)


def _default_target(week):
    """Muc tieu ROM mac dinh theo tuan phuc hoi."""
    if week <= 2:
        return 80
    if week <= 6:
        return 105
    return 125


@app.route("/game")
def game():
    """Game phuc hoi: gap dung muc tieu moi sut duoc.
    Muc tieu = ROM tuy chinh cua benh nhan (neu bac si/BN da dat), khong thi theo tuan."""
    if "user_id" not in session or session["role"] != "patient":
        return redirect(url_for("home"))
    conn = connect_db()
    cur = conn.cursor()
    cur.execute("SELECT week, rom_goal, rom_goal_days, rom_goal_start FROM users WHERE id = ?",
                (session["user_id"],))
    row = cur.fetchone()
    # ROM tot nhat tu truoc den nay (de hien tien do)
    cur.execute("SELECT MAX(angle) FROM measurements WHERE patient_id = ?", (session["user_id"],))
    bestrow = cur.fetchone()
    conn.close()

    week = row[0] if row and row[0] else 1
    goal, goal_days, goal_start = (row[1], row[2], row[3]) if row else (None, None, None)
    best_ever = int(bestrow[0]) if bestrow and bestrow[0] else 0

    # con bao nhieu ngay den han
    days_left = None
    if goal and goal_days and goal_start:
        try:
            start = datetime.datetime.strptime(goal_start, "%Y-%m-%d")
            days_left = goal_days - (datetime.datetime.now() - start).days
        except ValueError:
            days_left = None

    target = goal if goal else _default_target(week)
    return render_template("game.html", name=session["name"], week=week, target=target, tol=12,
                           goal=goal or 0, goal_days=goal_days or 14,
                           days_left=(days_left if days_left is not None else -999),
                           best_ever=best_ever, default_target=_default_target(week))


@app.route("/set_goal", methods=["POST"])
def set_goal():
    """Dat muc tieu ROM tuy chinh: dat <goal> do trong <days> ngay."""
    if "user_id" not in session:
        return jsonify({"success": False, "message": "Chua dang nhap"}), 401
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(success=False, message="Dữ liệu phải là đối tượng JSON."), 400
    try:
        goal = int(data.get("goal", 0))
        days = int(data.get("days", 14))
    except (TypeError, ValueError, OverflowError):
        return jsonify({"success": False, "message": "Gia tri khong hop le"}), 400
    if goal < 0 or goal > 150 or (goal > 0 and goal < 30) or not 1 <= days <= 180:
        return jsonify(success=False, message="Mục tiêu phải từ 30–150 độ, thời gian từ 1–180 ngày."), 400
    conn = connect_db()
    cur = conn.cursor()
    if goal <= 0:  # bo muc tieu -> quay ve mac dinh theo tuan
        cur.execute("UPDATE users SET rom_goal=NULL, rom_goal_days=NULL, rom_goal_start=NULL WHERE id=?",
                    (session["user_id"],))
        conn.commit()
        conn.close()
        return jsonify({"success": True, "cleared": True})
    goal = max(30, min(150, goal))
    days = max(1, min(180, days))
    start = datetime.datetime.now().strftime("%Y-%m-%d")
    cur.execute("UPDATE users SET rom_goal=?, rom_goal_days=?, rom_goal_start=? WHERE id=?",
                (goal, days, start, session["user_id"]))
    conn.commit()
    conn.close()
    return jsonify({"success": True, "goal": goal, "days": days, "start": start})


@app.route("/game_save", methods=["POST"])
def game_save():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(success=False, message="Dữ liệu phải là đối tượng JSON."), 400
    try:
        raw_rom = float(data["rom"])
        if isinstance(data["rom"], bool) or not math.isfinite(raw_rom) or not 0 <= raw_rom <= 180:
            raise ValueError
        rom = round(raw_rom, 1)
    except (TypeError, ValueError, OverflowError, KeyError):
        return jsonify(success=False, message="Góc ROM phải là số từ 0 đến 180."), 400
    source = data.get("source", "camera")
    if source not in ("camera", "imu"):
        return jsonify(success=False, message="Nguồn không hợp lệ."), 400
    with _state_lock:
        sampled_at = camera_last_data if source == "camera" else imu_last_data
        observed_max = max_rom_session if source == "camera" else imu_max_rom
        if not sampled_at or time.monotonic() - sampled_at > 3 or (source == "imu" and not imu_connected):
            return jsonify(success=False, message="Chưa có tín hiệu đo mới từ thiết bị."), 409
        if rom > observed_max + 1:
            return jsonify(success=False, message="Kết quả vượt quá góc ROM thiết bị ghi nhận."), 400
        conn = connect_db()
        try:
            ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            conn.execute("INSERT INTO measurements(patient_id, angle, time, source) VALUES (?, ?, ?, ?)",
                         (session["user_id"], rom, ts, source))
            conn.commit()
        finally:
            conn.close()
    return jsonify(success=True, rom=rom)


@app.route("/save_rom/<source>", methods=["POST"])
def save_rom(source):
    global max_rom_session, imu_max_rom, camera_last_data, imu_last_data
    if source not in ("camera", "imu"):
        return jsonify(success=False, message="Nguồn không hợp lệ."), 400
    with _state_lock:
        sampled_at = camera_last_data if source == "camera" else imu_last_data
        if not sampled_at or time.monotonic() - sampled_at > 3 or (source == "imu" and not imu_connected):
            return jsonify(success=False, message="Chưa có tín hiệu đo mới từ thiết bị."), 409
        value = max_rom_session if source == "camera" else imu_max_rom
        conn = connect_db()
        try:
            full_time_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cursor = conn.execute("INSERT INTO measurements(patient_id, angle, time, source) VALUES (?, ?, ?, ?)",
                                  (session["user_id"], value, full_time_str, source))
            conn.commit()
            new_id = cursor.lastrowid
        finally:
            conn.close()
        if source == "camera":
            max_rom_session = 0
            camera_last_data = 0.0
        else:
            imu_max_rom = 0
            imu_last_data = 0.0
    return jsonify(success=True, saved_rom=value, source=source, time=full_time_str, id=new_id)


@app.route("/delete_measurement/<int:mid>", methods=["POST"])
def delete_measurement(mid):
    """Xoa 1 luot do sai. Bac si xoa bat ky; benh nhan chi xoa cua minh."""
    if "user_id" not in session: return jsonify({"success": False, "message": "Chưa đăng nhập"}), 401
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT patient_id FROM measurements WHERE id = ?", (mid,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify({"success": False, "message": "Không tìm thấy lượt đo"})
    if session["role"] != "doctor" and session["user_id"] != row[0]:
        conn.close()
        return jsonify({"success": False, "message": "Không có quyền xóa"}), 403
    cursor.execute("DELETE FROM measurements WHERE id = ?", (mid,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/export/<int:pid>")
def export_excel(pid):
    """Xuat lich su do cua 1 benh nhan ra Excel. Bac si bat ky; benh nhan chi cua minh."""
    if "user_id" not in session: return redirect(url_for("home"))
    if session["role"] != "doctor" and session["user_id"] != pid:
        return "Không có quyền truy cập", 403
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT name, age, week FROM users WHERE id = ?", (pid,))
    u = cursor.fetchone()
    if not u:
        conn.close()
        return "Không tìm thấy bệnh nhân", 404
    name, age, week = u
    cursor.execute("SELECT angle, time, source FROM measurements WHERE patient_id = ? ORDER BY id ASC", (pid,))
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook(); ws = wb.active; ws.title = "Lich su do"
    hdr_fill = PatternFill("solid", fgColor="2F5597")
    hdr_font = Font(bold=True, color="FFFFFF")
    ws.append([f"BỆNH NHÂN: {name}  |  Tuổi: {age}  |  Tuần phục hồi: {week}"])
    ws.append(["STT", "Thời gian", "Góc ROM (độ)", "Nguồn đo"])
    for i, (ang, t, src) in enumerate(rows, 1):
        ws.append([i, t, ang, "Cảm biến IMU" if src == "imu" else "Camera"])
    for c in range(1, 5):
        cell = ws.cell(2, c); cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for col, w in zip("ABCD", [6, 22, 14, 16]): ws.column_dimensions[col].width = w

    safe = "".join(ch for ch in name if ch.isalnum() or ch in " _-").strip().replace(" ", "_")
    return _save_and_send(wb, f"lichsu_{safe or pid}.xlsx")

@app.route("/export_me")
def export_me():
    """Benh nhan tu xuat lich su cua chinh minh."""
    if "user_id" not in session or session["role"] != "patient":
        return redirect(url_for("home"))
    return export_excel(session["user_id"])

@app.route("/export_all")
def export_all():
    """Bac si xuat toan bo lich su cua moi benh nhan vao 1 file."""
    if "user_id" not in session or session["role"] != "doctor":
        return redirect(url_for("home"))
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT users.name, measurements.angle, measurements.time, measurements.source
        FROM measurements LEFT JOIN users ON users.id = measurements.patient_id
        ORDER BY users.name, measurements.id
    """)
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook(); ws = wb.active; ws.title = "Tat ca benh nhan"
    hdr_fill = PatternFill("solid", fgColor="2F5597")
    hdr_font = Font(bold=True, color="FFFFFF")
    ws.append(["Bệnh nhân", "Góc ROM (độ)", "Thời gian", "Nguồn đo"])
    for c in range(1, 5):
        cell = ws.cell(1, c); cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for name, ang, t, src in rows:
        ws.append([name, ang, t, "Cảm biến IMU" if src == "imu" else "Camera"])
    for col, w in zip("ABCD", [22, 14, 22, 16]): ws.column_dimensions[col].width = w

    return _save_and_send(wb, "lichsu_tat_ca_benh_nhan.xlsx")

@app.route("/delete_all_me", methods=["POST"])
def delete_all_me():
    """Benh nhan xoa TAT CA luot do cua chinh minh."""
    if "user_id" not in session or session["role"] != "patient":
        return jsonify({"success": False, "message": "Khong co quyen"}), 401
    conn = connect_db(); cur = conn.cursor()
    cur.execute("DELETE FROM measurements WHERE patient_id = ?", (session["user_id"],))
    conn.commit(); conn.close()
    return jsonify({"success": True})

@app.route("/delete_all/<int:pid>", methods=["POST"])
def delete_all(pid):
    """Xoa TAT CA luot do cua 1 benh nhan. Bac si bat ky; benh nhan chi cua minh."""
    if "user_id" not in session: return jsonify({"success": False}), 401
    if session["role"] != "doctor" and session["user_id"] != pid:
        return jsonify({"success": False, "message": "Khong co quyen"}), 403
    conn = connect_db(); cur = conn.cursor()
    cur.execute("DELETE FROM measurements WHERE patient_id = ?", (pid,))
    conn.commit(); conn.close()
    return jsonify({"success": True})

@app.route("/export_compare/<int:pid>")
def export_compare(pid):
    """Xuat Excel SO SANH Camera vs Cam bien (2 duong tren 1 bieu do)."""
    if "user_id" not in session: return redirect(url_for("home"))
    if session["role"] != "doctor" and session["user_id"] != pid:
        return "Không có quyền truy cập", 403
    conn = connect_db(); cur = conn.cursor()
    cur.execute("SELECT name FROM users WHERE id = ?", (pid,))
    u = cur.fetchone()
    if not u:
        conn.close(); return "Không tìm thấy bệnh nhân", 404
    name = u[0]
    cur.execute("SELECT angle, time FROM measurements WHERE patient_id=? AND source='camera' ORDER BY id", (pid,))
    cams = cur.fetchall()
    cur.execute("SELECT angle, time FROM measurements WHERE patient_id=? AND source='imu' ORDER BY id", (pid,))
    imus = cur.fetchall()
    conn.close()

    wb = Workbook(); ws = wb.active; ws.title = "So sanh"
    hdr_fill = PatternFill("solid", fgColor="0E82FD"); hdr_font = Font(bold=True, color="FFFFFF")
    ws.append([f"SO SÁNH Camera vs Cảm biến — Bệnh nhân: {name}"])
    ws.append(["Lượt đo", "Thời gian", "Camera (độ)", "Cảm biến IMU (độ)", "Chênh lệch (độ)"])
    n = max(len(cams), len(imus))
    for i in range(n):
        cval = cams[i][0] if i < len(cams) else None
        ival = imus[i][0] if i < len(imus) else None
        tval = cams[i][1] if i < len(cams) else (imus[i][1] if i < len(imus) else "")
        diff = abs(cval - ival) if (cval is not None and ival is not None) else None
        ws.append([i + 1, tval, cval, ival, diff])
    last = n + 2
    for col_i in range(1, 6):
        cell = ws.cell(2, col_i); cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for col_l, wdt in zip("ABCDE", [9, 20, 14, 16, 14]): ws.column_dimensions[col_l].width = wdt

    if n > 0:
        chart = LineChart(); chart.title = "Camera vs Cảm biến (góc ROM)"
        chart.y_axis.title = "Góc (độ)"; chart.x_axis.title = "Lượt đo"
        chart.height = 10; chart.width = 22
        chart.add_data(Reference(ws, min_col=3, max_col=4, min_row=2, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=last))
        # Hien ro CA 2 TRUC + thang doc len 180 do
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 180
        chart.y_axis.majorUnit = 30
        chart.x_axis.majorTickMark = "out"
        chart.y_axis.majorTickMark = "out"
        if len(chart.series) >= 1:
            s0 = chart.series[0]; s0.smooth = True
            s0.graphicalProperties = GraphicalProperties(); s0.graphicalProperties.line = LineProperties(solidFill="0E82FD", w=26000)
        if len(chart.series) >= 2:
            s1 = chart.series[1]; s1.smooth = True
            s1.graphicalProperties = GraphicalProperties(); s1.graphicalProperties.line = LineProperties(solidFill="16A34A", w=26000)
        ws.add_chart(chart, "G2")

    # ===== ĐÁNH GIÁ ĐỒNG THUẬN 2 PHƯƠNG PHÁP (Camera vs Cảm biến) =====
    paired = [(cams[i][0], imus[i][0]) for i in range(min(len(cams), len(imus)))
              if cams[i][0] is not None and imus[i][0] is not None]
    if paired:
        diffs = [c - m for c, m in paired]          # sai so co dau (camera - cam bien)
        ad = [abs(d) for d in diffs]
        nP = len(paired)
        mae = sum(ad) / nP                           # sai so tuyet doi trung binh
        rmse = (sum(d * d for d in diffs) / nP) ** 0.5
        bias = sum(diffs) / nP                        # do lech he thong (>0: camera doc cao hon)
        mx = max(ad)
        cam_v = [p[0] for p in paired]; imu_v = [p[1] for p in paired]
        mc = sum(cam_v) / nP; mi = sum(imu_v) / nP
        num = sum((cam_v[k] - mc) * (imu_v[k] - mi) for k in range(nP))
        den = (sum((x - mc) ** 2 for x in cam_v) * sum((x - mi) ** 2 for x in imu_v)) ** 0.5
        corr = (num / den) if den > 0 else 0.0
        if mae <= 5:    verdict = "Đồng thuận TỐT (sai số TB ≤ 5°) — 2 phương pháp tin cậy như nhau"
        elif mae <= 10: verdict = "Đồng thuận KHÁ (sai số 5–10°)"
        else:           verdict = "Chênh lệch LỚN (>10°) — nên kiểm tra lại đặt cảm biến/khung camera"
        who = "Camera đo cao hơn" if bias > 1 else ("Cảm biến đo cao hơn" if bias < -1 else "Gần như cân bằng")

        r0 = last + 2
        title_cell = ws.cell(r0, 1, "ĐÁNH GIÁ ĐỒNG THUẬN 2 PHƯƠNG PHÁP")
        title_cell.font = Font(bold=True, color="FFFFFF"); title_cell.fill = PatternFill("solid", fgColor="0E82FD")
        ws.cell(r0, 2).fill = PatternFill("solid", fgColor="0E82FD")
        rows_stat = [
            ("Số buổi so sánh", nP),
            ("Sai số tuyệt đối trung bình - MAE (độ)", round(mae, 1)),
            ("Sai số quân phương - RMSE (độ)", round(rmse, 1)),
            ("Độ lệch hệ thống - bias (độ)", f"{round(bias, 1)}  ({who})"),
            ("Sai số lớn nhất (độ)", round(mx, 1)),
            ("Hệ số tương quan (1 = trùng khớp)", round(corr, 3)),
            ("Kết luận", verdict),
        ]
        for k, (lab, val) in enumerate(rows_stat):
            ws.cell(r0 + 1 + k, 1, lab).font = Font(bold=True)
            ws.cell(r0 + 1 + k, 2, val)

    safe = "".join(ch for ch in name if ch.isalnum() or ch in " _-").strip().replace(" ", "_")
    return _save_and_send(wb, f"sosanh_{safe or pid}.xlsx")

@app.route("/export_compare_me")
def export_compare_me():
    """Benh nhan tu xuat Excel so sanh cua chinh minh."""
    if "user_id" not in session or session["role"] != "patient":
        return redirect(url_for("home"))
    return export_compare(session["user_id"])



@app.route("/admin")
def admin():
    if "user_id" not in session or session["role"] != "doctor":
        return redirect(url_for("home"))

    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM users WHERE role = 'patient'")
    total_patients = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM measurements")
    total_measurements = cursor.fetchone()[0]
    cursor.execute("SELECT id, name, age, week, username FROM users WHERE role = 'patient'")
    patients = cursor.fetchall()
    cursor.execute("""
        SELECT measurements.id, users.name, measurements.angle, measurements.time
        FROM measurements
        LEFT JOIN users ON users.id = measurements.patient_id
        ORDER BY measurements.id DESC
    """)
    measurements = cursor.fetchall()
    conn.close()

    return render_template("admin_dashboard.html", total_patients=total_patients, total_measurements=total_measurements, patients=patients, measurements=measurements)

@app.route("/admin/patient/<int:pid>")
def admin_patient_view(pid):
    if "user_id" not in session or session["role"] != "doctor":
        return redirect(url_for("home"))
        
    conn = connect_db()
    cursor = conn.cursor()
    
    cursor.execute("SELECT id, name, age, week, username, doctor_advice FROM users WHERE id=?", (pid,))
    user = cursor.fetchone()
    
    if not user:
        conn.close()
        return "Không tìm thấy thông tin bệnh nhân này!", 404
        
    cursor.execute("SELECT angle, time, source, id FROM measurements WHERE patient_id=? ORDER BY id DESC", (pid,))
    history = cursor.fetchall()
    conn.close()

    return render_template("admin_patients_view.html", user=user, history=history)


@app.route("/admin/save_advice/<int:pid>", methods=["POST"])
def admin_save_advice(pid):
    if "user_id" not in session or session["role"] != "doctor":
        return jsonify({"success": False, "message": "Quyền truy cập bị từ chối"}), 403
    
    advice_text = request.form.get("doctor_advice", "").strip()
    if len(advice_text) > 5000:
        return jsonify(success=False, message="Lời khuyên quá dài (tối đa 5000 ký tự)."), 400
    
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE users SET doctor_advice = ? WHERE id = ?", (advice_text, pid))
    conn.commit()
    conn.close()
    
    return redirect(url_for("admin_patient_view", pid=pid))

def main():
    logging.basicConfig(level=logging.INFO)
    init_db()
    _initialized_databases.add(app.config["DATABASE"])
    try:
        app.run(host=os.getenv("ACL_HOST", "127.0.0.1"),
                port=_env_int("ACL_PORT", 5000, 1, 65535),
                debug=_env_bool("ACL_DEBUG"), threaded=True, use_reloader=False)
    finally:
        shutdown_resources()


if __name__ == "__main__":
    main()
